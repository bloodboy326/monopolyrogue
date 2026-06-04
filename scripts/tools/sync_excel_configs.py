from __future__ import annotations

import json
from copy import deepcopy
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[2]
CARD_XLSX = Path(r"C:\Users\yu\Desktop\card.xlsx")
RELIC_XLSX = Path(r"C:\Users\yu\Documents\xwechat_files\wxid_8dzrcbxqxc0322_5fea\msg\file\2026-06\relics.xlsx")
MONSTER_XLSX = Path(r"C:\Users\yu\Desktop\monster_new.xlsx")
TILE_CONFIG = ROOT / "data" / "tile_config.json"
RELIC_CONFIG = ROOT / "data" / "relic_config.json"
MONSTER_CONFIG = ROOT / "data" / "monster_config.json"
MAP_CONFIG = ROOT / "data" / "map_config.json"


def tid(value: int) -> str:
    return f"T{value:03d}"


def rid(value: int) -> str:
    return f"R{value:03d}"


def read_rows(path: Path) -> list[list[Any]]:
    workbook = load_workbook(path, data_only=True)
    sheet = workbook.active
    return [
        list(row)
        for row in sheet.iter_rows(min_row=2, values_only=True)
        if any(value is not None for value in row)
    ]


def text(value: Any, fallback: str = "") -> str:
    if value is None:
        return fallback
    return str(value).strip()


def number(value: Any, fallback: int = 0) -> int:
    if value is None or value == "":
        return fallback
    return int(value)


def flag(value: Any) -> bool:
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return int(value) != 0
    return text(value).lower() in {"1", "true", "yes", "y", "是", "有"}


def write_json(path: Path, data: dict[str, Any]) -> None:
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def sort_tile(tile: dict[str, Any]) -> tuple[int, str]:
    tile_id = str(tile.get("id", ""))
    if tile_id.startswith("T") and tile_id[1:].isdigit():
        return int(tile_id[1:]), tile_id
    return 9999, tile_id


def rarity_code(rarity: str) -> str:
    return {
        "基础牌": "basic",
        "普通": "common",
        "稀有": "rare",
        "非凡": "uncommon",
        "诅咒": "curse",
        "boss": "boss",
    }.get(rarity, "common")


def _legacy_sync_tiles() -> None:
    config = json.loads(TILE_CONFIG.read_text(encoding="utf-8"))
    tiles = {str(tile["id"]): tile for tile in config.get("tiles", [])}
    card_rows = {number(row[0]): row for row in read_rows(CARD_XLSX) if row[0] is not None}

    extra_defs: dict[int, dict[str, Any]] = {
        14: {
            "tags": ["skill", "energy", "void"],
            "icon_kind": "charge",
            "color": [0.15, 0.74, 0.65, 1.0],
            "accent": [1.0, 0.86, 0.22, 1.0],
            "durability": 2,
            "effects": [
                {"type": "add_rolls", "value": 3},
                {"type": "generate_tile", "tile": "T072", "count": 1},
            ],
            "weakEffects": [
                {"type": "add_rolls", "value": 1},
                {"type": "generate_tile", "tile": "T072", "count": 1},
            ],
        },
        27: {
            "tags": ["skill", "warp", "energy"],
            "icon_kind": "lab",
            "effects": [{"type": "add_rolls", "value": 3, "condition": {"reason": "warp"}}],
        },
        28: {
            "tags": ["skill", "warp", "attack"],
            "icon_kind": "charge",
            "effects": [{"type": "damage", "value": 3}],
            "passEffects": [{"type": "damage", "value": 4, "condition": {"reason_in": ["warp_passed"]}}],
        },
        29: {
            "tags": ["skill", "warp", "defense"],
            "icon_kind": "combo",
            "effects": [{"type": "block", "value": 5}],
            "passEffects": [{"type": "block", "value": 4, "condition": {"reason_in": ["warp_passed"]}}],
        },
        40: {
            "tags": ["skill", "void"],
            "icon_kind": "void",
            "color": [0.36, 0.22, 0.62, 1.0],
            "accent": [0.88, 0.52, 1.0, 1.0],
            "durability": 3,
            "effects": [
                {"type": "increment_counter", "scope": "turn", "counter": "player_damage_cap_1", "value": 1},
                {"type": "generate_tile", "tile": "T072", "count": 1},
            ],
            "weakEffects": [],
        },
        45: {
            "tags": ["skill", "defense"],
            "icon_kind": "shield_wall",
            "effects": [
                {"type": "increment_counter", "scope": "battle", "counter": "roll_block_bonus", "value": 2},
                {"type": "destroy_self"},
            ],
        },
        49: {
            "tags": ["skill", "durability"],
            "icon_kind": "power",
            "color": [0.22, 0.52, 0.90, 1.0],
            "accent": [0.66, 0.92, 1.0, 1.0],
            "effects": [
                {"type": "block", "value": 8},
                {"type": "add_durability_all", "value": 1},
            ],
        },
        55: {
            "tags": ["skill", "energy"],
            "icon_kind": "coin",
            "color": [0.15, 0.74, 0.65, 1.0],
            "accent": [1.0, 0.86, 0.22, 1.0],
            "effects": [
                {"type": "add_rolls", "value": 2},
                {"type": "destroy_self"},
            ],
        },
        63: {
            "tags": ["attack"],
            "icon_kind": "sword",
            "color": [0.96, 0.30, 0.36, 1.0],
            "accent": [0.40, 1.0, 0.96, 1.0],
            "durability": 2,
            "effects": [{"type": "damage", "value": 28}],
            "weakEffects": [{"type": "damage", "value": 10}],
        },
        66: {
            "tags": ["skill", "heal"],
            "icon_kind": "heal",
            "color": [0.15, 0.74, 0.42, 1.0],
            "accent": [1.0, 0.92, 0.28, 1.0],
            "effects": [{"type": "heal_player", "value": 7}, {"type": "destroy_self"}],
        },
        67: {
            "tags": ["curse", "sticky"],
            "icon_kind": "slime_mud",
            "color": [0.34, 0.84, 0.30, 1.0],
            "accent": [0.9, 1.0, 0.3, 1.0],
            "passEffects": [{"type": "stop_movement"}],
            "effects": [{"type": "destroy_self"}],
        },
        68: {
            "tags": ["curse", "slime", "mud"],
            "icon_kind": "slime_mud",
            "color": [0.20, 0.46, 0.28, 1.0],
            "accent": [0.60, 0.95, 0.38, 1.0],
            "effects": [{"type": "destroy_self"}],
        },
        69: {
            "tags": ["curse", "bone"],
            "icon_kind": "bone",
            "color": [0.58, 0.56, 0.50, 1.0],
            "accent": [0.95, 0.92, 0.75, 1.0],
            "passEffects": [
                {"type": "destroy_self", "condition": {"pass_through_only": True}},
            ],
            "effects": [{"type": "destroy_self"}],
        },
        70: {
            "tags": ["curse", "moon"],
            "icon_kind": "moon",
            "color": [0.20, 0.20, 0.40, 1.0],
            "accent": [0.78, 0.88, 1.0, 1.0],
            "counters": {"moon": 0},
            "eventHooks": ["moon_counter"],
            "effects": [],
        },
        71: {
            "tags": ["curse", "arrow"],
            "icon_kind": "arrow",
            "color": [0.78, 0.20, 0.30, 1.0],
            "accent": [1.0, 0.78, 0.28, 1.0],
            "effects": [
                {"type": "damage_player", "value": 5, "piercing": True},
                {"type": "destroy_self"},
            ],
        },
        72: {
            "tags": ["curse", "void"],
            "icon_kind": "void",
            "color": [0.28, 0.18, 0.46, 1.0],
            "accent": [0.86, 0.46, 1.0, 1.0],
            "effects": [],
        },
        73: {
            "tags": ["attack", "aoe"],
            "icon_kind": "sword",
            "color": [0.94, 0.30, 0.36, 1.0],
            "accent": [1.0, 0.86, 0.22, 1.0],
            "effects": [{"type": "damage_all", "value": 8}],
        },
        74: {
            "tags": ["attack", "aoe"],
            "icon_kind": "sword",
            "color": [0.94, 0.30, 0.36, 1.0],
            "accent": [0.44, 1.0, 0.94, 1.0],
            "durability": 2,
            "effects": [{"type": "damage_all", "value": 16}],
            "weakEffects": [{"type": "damage_all", "value": 8}],
        },
        75: {
            "tags": ["skill", "energy", "destroy"],
            "icon_kind": "charge",
            "durability": 2,
            "effects": [
                {"type": "add_rolls", "value": 1},
                {"type": "destroy_tiles", "targetRule": {"type": "random", "count": 1, "excludeCurrent": True}, "destroyMode": {"type": "configured"}},
            ],
            "weakEffects": [],
        },
        76: {
            "tags": ["skill", "supply"],
            "icon_kind": "lab",
            "effects": [
                {"type": "add_converge_to_tiles", "count": 2},
                {"type": "destroy_self"},
            ],
        },
        77: {
            "tags": ["skill", "aoe"],
            "icon_kind": "sword",
            "effects": [
                {"type": "increment_counter", "scope": "battle", "counter": "stray_bullet_storm", "value": 1},
                {"type": "destroy_self"},
            ],
        },
        78: {
            "tags": ["skill", "durability"],
            "icon_kind": "charge",
            "durability": 1,
            "effects": [
                {"type": "increment_counter", "scope": "turn", "counter": "no_durability_cost_turn", "value": 1},
                {"type": "set_tile_weak"},
            ],
            "weakEffects": [],
        },
        79: {
            "tags": ["curse", "weakness"],
            "icon_kind": "void",
            "effects": [{"type": "destroy_self"}],
        },
        80: {
            "tags": ["skill", "gold"],
            "icon_kind": "coin",
            "color": [0.95, 0.60, 0.18, 1.0],
            "accent": [1.0, 0.95, 0.28, 1.0],
            "effects": [
                {"type": "add_coins", "value": 15},
                {"type": "destroy_self"},
            ],
        },
        81: {
            "tags": ["curse", "headache"],
            "icon_kind": "void",
            "effects": [],
        },
        82: {
            "tags": ["skill", "basic", "repair"],
            "icon_kind": "repair",
            "color": [0.18, 0.62, 0.78, 1.0],
            "accent": [1.0, 0.86, 0.22, 1.0],
            "effects": [
                {"type": "block", "value": 4},
                {"type": "add_durability_all", "value": 1},
            ],
        },
        83: {
            "tags": ["attack", "basic", "energy"],
            "icon_kind": "charge",
            "color": [0.92, 0.24, 0.32, 1.0],
            "accent": [0.52, 0.92, 1.0, 1.0],
            "effects": [
                {"type": "damage", "value": 4},
                {"type": "add_rolls", "value": 1},
            ],
        },
    }

    changed_ids = {
        number(row[0])
        for row in read_rows(CARD_XLSX)
        if row[0] is not None and len(row) > 8 and flag(row[8])
    }
    sync_ids = sorted({14, 27, 28, 29, 66, 67, 68, 69, 70, 71, 72, 82, 83}.union(changed_ids).intersection(card_rows.keys()))
    for numeric_id in sync_ids:
        row = card_rows[numeric_id]
        tile_id = tid(numeric_id)
        current = tiles.get(tile_id, {"id": tile_id})
        rarity = text(row[3])
        current.update(
            {
                "id": tile_id,
                "type": text(row[1]),
                "name": text(row[2]),
                "rarity": rarity,
                "rarity_code": rarity_code(rarity),
                "description": text(row[4]),
                "displayDescription": text(row[5], text(row[4])),
                "selectable": flag(row[6]),
                "temporary": flag(row[7]),
                "destroy_after_battle": flag(row[7]),
                "changed": flag(row[8] if len(row) > 8 else 0),
            }
        )
        current.update(extra_defs.get(numeric_id, {}))
        tiles[tile_id] = current

    config["tiles"] = sorted(tiles.values(), key=sort_tile)
    config["start_tiles"] = ["T001", "T002", "T001", "T002", "T001", "T002", "T003", "T004"]
    write_json(TILE_CONFIG, config)


def sync_tiles() -> None:
    config = json.loads(TILE_CONFIG.read_text(encoding="utf-8"))
    existing_tiles = config.get("tiles", [])
    existing_by_id = {str(tile.get("id", "")): tile for tile in existing_tiles if isinstance(tile, dict)}
    existing_by_name = {
        text(tile.get("name", "")): tile
        for tile in existing_tiles
        if isinstance(tile, dict) and text(tile.get("name", ""))
    }
    card_rows = {number(row[0]): row for row in read_rows(CARD_XLSX) if row[0] is not None}

    def tile_id_by_name(name: str, fallback: str) -> str:
        for numeric_id, row in card_rows.items():
            if text(row[2]) == name:
                return tid(numeric_id)
        return fallback

    overrides_by_name: dict[str, dict[str, Any]] = {
        "修修补补": {
            "tags": ["skill", "basic", "repair", "durability"],
            "icon_kind": "repair",
            "color": [0.22, 0.52, 0.90, 1.0],
            "accent": [0.66, 0.92, 1.0, 1.0],
            "effects": [
                {"type": "block", "value": 8},
                {"type": "add_durability_all", "value": 1},
            ],
        },
        "充能剑": {
            "tags": ["attack", "basic", "energy", "durability"],
            "icon_kind": "charge_sword",
            "color": [0.92, 0.25, 0.32, 1.0],
            "accent": [1.0, 0.82, 0.22, 1.0],
            "durability": 2,
            "effects": [{"type": "damage", "value": 10}],
            "weakEffects": [{"type": "damage", "value": 2}],
        },
    }

    table_ids = set(card_rows.keys())
    preserved: list[dict[str, Any]] = []
    for tile in existing_tiles:
        if not isinstance(tile, dict):
            continue
        tile_id = str(tile.get("id", ""))
        numeric = int(tile_id[1:]) if tile_id.startswith("T") and tile_id[1:].isdigit() else -1
        if tile_id == "T000" or numeric >= 900 or numeric < 0:
            preserved.append(deepcopy(tile))

    synced: list[dict[str, Any]] = []
    for numeric_id, row in sorted(card_rows.items()):
        tile_id = tid(numeric_id)
        name = text(row[2], tile_id)
        current = deepcopy(existing_by_name.get(name, existing_by_id.get(tile_id, {"id": tile_id})))
        rarity = text(row[3])
        current.update(
            {
                "id": tile_id,
                "type": text(row[1]),
                "name": name,
                "rarity": rarity,
                "rarity_code": rarity_code(rarity),
                "description": text(row[4]),
                "displayDescription": text(row[5], text(row[4])),
                "selectable": flag(row[6]),
                "temporary": flag(row[7]),
                "destroy_after_battle": flag(row[7]),
                "changed": flag(row[8] if len(row) > 8 else 0),
                "tile_name": name,
                "tile_rare": rarity,
                "tile_describe": text(row[5], text(row[4])),
            }
        )
        current.update(deepcopy(overrides_by_name.get(name, {})))
        synced.append(current)

    config["tiles"] = sorted(preserved + synced, key=sort_tile)
    attack = tile_id_by_name("普通攻击", "T001")
    defense = tile_id_by_name("普通防御", "T002")
    config["start_tiles"] = [
        attack,
        defense,
        attack,
        defense,
        attack,
        defense,
        tile_id_by_name("修修补补", "T003"),
        tile_id_by_name("充能剑", "T004"),
    ]
    write_json(TILE_CONFIG, config)


def tile_id_from_config_name(name: str, fallback: str = "T000") -> str:
    if not TILE_CONFIG.exists():
        return fallback
    config = json.loads(TILE_CONFIG.read_text(encoding="utf-8"))
    for tile in config.get("tiles", []):
        if isinstance(tile, dict) and text(tile.get("name", "")) == name:
            return str(tile.get("id", fallback))
    return fallback


def sync_relics() -> None:
    gold_tile = tile_id_from_config_name("黄金", "T081")
    headache_tile = tile_id_from_config_name("头疼", "T082")
    effect_by_index: dict[int, list[dict[str, Any]]] = {
        1: [{"type": "dice_roll_energy", "threshold": 10, "amount": 1}],
        2: [{"type": "battle_start_tile", "tile": gold_tile, "count": 1}],
        3: [{"type": "battle_start_strength", "amount": 1}],
        4: [{"type": "battle_start_dexterity", "amount": 1}],
        5: [{"type": "turn_start_rolls", "turn": 1, "amount": 1}],
        6: [{"type": "turn_interval_rolls", "interval": 4, "amount": 1}],
        7: [{"type": "turn_start_block", "turn": 1, "amount": 10}],
        8: [{"type": "turn_start_block", "turn": 2, "amount": 13}],
        9: [{"type": "turn_start_block", "turn": 3, "amount": 18}],
        10: [{"type": "durability_interval_recharge", "threshold": 10, "amount": 1}],
        11: [{"type": "warp_interval_rolls", "threshold": 3, "amount": 1}],
        12: [{"type": "attack_interval_strength", "threshold": 3, "amount": 1}],
        13: [
            {"type": "turn_start_rolls", "amount": 1},
            {"type": "dice_roll_hp_loss", "threshold": 10, "amount": 1},
        ],
        14: [
            {"type": "turn_start_rolls", "amount": 1},
            {"type": "battle_start_tile", "tile": headache_tile, "count": 2},
        ],
        15: [],
    }
    icon_kinds = {
        1: "dice",
        2: "gold_eye",
        3: "strength",
        4: "dexterity",
        5: "orb",
        6: "chip",
        7: "white_robe",
        8: "red_robe",
        9: "purple_robe",
        10: "support",
        11: "wormhole",
        12: "target",
        13: "blood_dice",
        14: "potion",
        15: "portal",
    }
    relics: list[dict[str, Any]] = []
    for index, row in enumerate(read_rows(RELIC_XLSX), start=1):
        rarity = text(row[3], "普通")
        price_text = text(row[4])
        relics.append(
            {
                "id": rid(index),
                "name": text(row[1], f"遗物{index}"),
                "type": text(row[2], "普通"),
                "rarity": rarity,
                "rarity_code": rarity_code(rarity),
                "price": 0 if price_text == "/" else number(row[4], 0),
                "description": text(row[5]),
                "displayDescription": text(row[6], text(row[5])),
                "icon_kind": icon_kinds.get(index, "relic"),
                "effects": effect_by_index.get(index, []),
            }
        )
    write_json(RELIC_CONFIG, {"relics": relics})


MONSTER_META = {
    "goblin": (1, "goblin"),
    "slime": (2, "slime"),
    "skeleton_soldier": (3, "skeleton_soldier"),
    "murloc": (4, "murloc"),
    "skeleton_archer": (5, "skeleton_archer"),
    "frost_giant": (6, "frost_giant"),
    "baby_dragon": (9, "baby_dragon"),
    "skeleton_king": (10, "skeleton_king"),
    "werewolf": (11, "werewolf"),
    "headless_knight": (12, "headless_knight"),
    "medusa": (13, "medusa"),
    "void_eye": (14, "void_eye"),
    "dice_demon": (15, "dice_demon"),
}


def effect(intent_id: str, order: int, effect_type: str, value: int = 0, target: str = "", param: str = "", **extra: Any) -> dict[str, Any]:
    data: dict[str, Any] = {
        "intent_id": intent_id,
        "order": order,
        "effect_type": effect_type,
        "target": target,
        "value": value,
        "duration": int(extra.pop("duration", 0)),
        "param": param,
    }
    data.update(extra)
    return data


def intent(intent_id: str, intent_type: str, name: str, telegraph: str, description: str, threat: int = 1, **extra: Any) -> dict[str, Any]:
    role = extra.pop("role", "DAMAGE" if intent_type == "ATTACK" else "DEFENSE" if intent_type == "DEFEND" else "CONTROL")
    data = {
        "intent_id": intent_id,
        "intent_type": intent_type,
        "intent_role": role,
        "name": name,
        "telegraph": telegraph,
        "description": description,
        "threat_level": threat,
    }
    data.update(extra)
    return data


def pool(pool_id: str, intent_id: str, weight: int = 33, min_turn: int = 1, cooldown: int = 0, max_repeat: int = 2, require: str = "", forbid: str = "", max_uses: int = 0) -> dict[str, Any]:
    data = {
        "pool_id": pool_id,
        "intent_id": intent_id,
        "weight": weight,
        "min_turn": min_turn,
        "cooldown": cooldown,
        "max_repeat": max_repeat,
        "require": require,
        "forbid": forbid,
    }
    if max_uses > 0:
        data["max_uses"] = max_uses
    return data


def monster_from_row(monster_id: str, row: list[Any], art_key: str, passive: dict[str, Any] | None = None) -> dict[str, Any]:
    data: dict[str, Any] = {
        "monster_id": monster_id,
        "name": text(row[1]),
        "type": text(row[2]),
        "max_hp": number(row[3], 40),
        "region": text(row[4]),
        "start_phase": f"{monster_id}_p1",
        "art_key": art_key,
        "source_description": text(row[5]),
    }
    if passive:
        data["passive"] = passive
    return data


def group_monster(monster_id: str, name: str, monster_type: str, region: str, units: list[str], art_key: str) -> dict[str, Any]:
    return {
        "monster_id": monster_id,
        "name": name,
        "type": monster_type,
        "region": region,
        "max_hp": 1,
        "start_phase": "",
        "art_key": art_key,
        "units": [{"monster_id": unit_id} for unit_id in units],
    }


def sync_monsters() -> None:
    rows_by_id = {number(row[0]): row for row in read_rows(MONSTER_XLSX) if row[0] is not None}
    fire_tile = tile_id_from_config_name("烈火", "T065")
    stun_tile = tile_id_from_config_name("晕眩", "T066")
    snot_tile = tile_id_from_config_name("鼻涕液", "T068")
    mud_tile = tile_id_from_config_name("淤泥", "T069")
    bone_tile = tile_id_from_config_name("骨头", "T070")
    moon_tile = tile_id_from_config_name("月圆", "T071")
    split_arrow_tile = tile_id_from_config_name("分裂箭", "T072")
    void_tile = tile_id_from_config_name("虚空", "T073")

    monsters: list[dict[str, Any]] = []
    for monster_id, (source_id, art_key) in MONSTER_META.items():
        passive: dict[str, Any] | None = None
        if monster_id in {"skeleton_soldier", "skeleton_archer", "skeleton_king"}:
            passive = {"on_damaged_add_tile": bone_tile, "count": 1}
        elif monster_id == "dice_demon":
            passive = {"strength_per_player_rolls": 5, "strength_amount": 1}
        monsters.append(monster_from_row(monster_id, rows_by_id[source_id], art_key, passive))

    monsters.append(group_monster("skeleton_pair", "骷髅巡逻队", "群体小怪", "1层后期", ["skeleton_soldier", "skeleton_archer"], "skeleton_soldier"))
    guard_a = {"monster_id": "guard_a", "name": "卫兵A", "type": "群体单位", "region": "1层后期", "max_hp": 40, "start_phase": "guard_a_p1", "art_key": "guard_a"}
    guard_b = {"monster_id": "guard_b", "name": "卫兵B", "type": "群体单位", "region": "1层后期", "max_hp": 38, "start_phase": "guard_b_p1", "art_key": "guard_b"}
    monsters.append(guard_a)
    monsters.append(guard_b)
    monsters.append(group_monster("guard_pair", "卫兵双人组", "群体小怪", "1层后期", ["guard_a", "guard_b"], "guard_a"))

    phases: list[dict[str, Any]] = []
    for item in monsters:
        if item.get("units"):
            continue
        phase = {
            "phase_id": f"{item['monster_id']}_p1",
            "monster_id": item["monster_id"],
            "enter_condition": "START",
            "on_enter_effect": "",
            "intent_pool": f"{item['monster_id']}_pool",
        }
        if item["monster_id"] == "murloc":
            phase["sequence"] = ["murloc_swamp_charge", "murloc_dive"]
            phase["sequence_start_turn"] = 1
        elif item["monster_id"] == "frost_giant":
            phase["sequence"] = ["frost_giant_ice_shield", "frost_giant_frost_strike"]
            phase["sequence_start_turn"] = 2
        phases.append(phase)

    pools = [
        pool("goblin_pool", "goblin_snot", 100, require="TURN<=1", cooldown=99, max_repeat=1),
        pool("goblin_pool", "goblin_mace", 50, min_turn=2),
        pool("goblin_pool", "goblin_shield", 50, min_turn=2, cooldown=1, max_repeat=1),
        pool("slime_pool", "slime_attack", 40),
        pool("slime_pool", "slime_defend", 25, cooldown=1, max_repeat=1),
        pool("slime_pool", "slime_mud", 35, cooldown=2, max_repeat=1),
        pool("skeleton_soldier_pool", "skeleton_swing", 55),
        pool("skeleton_soldier_pool", "skeleton_spikes", 45, cooldown=1),
        pool("murloc_pool", "murloc_swamp_charge", 50),
        pool("murloc_pool", "murloc_dive", 50),
        pool("skeleton_archer_pool", "skeleton_archer_volley", 60),
        pool("skeleton_archer_pool", "skeleton_archer_empower", 40, cooldown=2, max_repeat=1),
        pool("frost_giant_pool", "frost_giant_ice_power", 100, require="TURN<=1", cooldown=99, max_repeat=1),
        pool("frost_giant_pool", "frost_giant_ice_shield", 50, min_turn=2),
        pool("frost_giant_pool", "frost_giant_frost_strike", 50, min_turn=2),
        pool("guard_a_pool", "guard_thrust", 34),
        pool("guard_a_pool", "guard_formation", 33, cooldown=1),
        pool("guard_a_pool", "guard_bash", 33),
        pool("guard_b_pool", "guard_thrust", 34),
        pool("guard_b_pool", "guard_formation", 33, cooldown=1),
        pool("guard_b_pool", "guard_bash", 33),
        pool("baby_dragon_pool", "baby_dragon_flame", 42),
        pool("baby_dragon_pool", "baby_dragon_burn", 28, cooldown=2, max_uses=2),
        pool("baby_dragon_pool", "baby_dragon_tail", 30),
        pool("skeleton_king_pool", "skeleton_king_bone_spikes", 1000, require=f"TILE_COUNT:{bone_tile}>=3"),
        pool("skeleton_king_pool", "skeleton_king_bone_shield", 35, forbid=f"TILE_COUNT:{bone_tile}>=3", cooldown=1),
        pool("skeleton_king_pool", "skeleton_king_slam", 65, forbid=f"TILE_COUNT:{bone_tile}>=3"),
        pool("werewolf_pool", "werewolf_moon", 100, require="TURN<=1", cooldown=99, max_repeat=1),
        pool("werewolf_pool", "werewolf_claw", 55, min_turn=2, require=f"TILE_COUNT:{moon_tile}>=1"),
        pool("werewolf_pool", "werewolf_lick", 45, min_turn=2, require=f"TILE_COUNT:{moon_tile}>=1", cooldown=1),
        pool("werewolf_pool", "werewolf_hysteria", 60, min_turn=2, require=f"TILE_COUNT:{moon_tile}<=0"),
        pool("werewolf_pool", "werewolf_moonbath", 40, min_turn=2, require=f"TILE_COUNT:{moon_tile}<=0", cooldown=1),
        pool("headless_knight_pool", "headless_swing", 38),
        pool("headless_knight_pool", "headless_flame_combo", 32, min_turn=2),
        pool("headless_knight_pool", "headless_wildfire", 30, cooldown=2, max_repeat=1),
        pool("medusa_pool", "medusa_split_arrow", 1000, require=f"TILE_COUNT:{split_arrow_tile}<=0"),
        pool("medusa_pool", "medusa_petrify", 34, forbid=f"TILE_COUNT:{split_arrow_tile}<=0"),
        pool("medusa_pool", "medusa_snake_bite", 66, forbid=f"TILE_COUNT:{split_arrow_tile}<=0"),
        pool("void_eye_pool", "void_eye_void_gaze", 1000, require="TURN_MOD:5=0"),
        pool("void_eye_pool", "void_eye_hot_ray", 55, forbid="TURN_MOD:5=0"),
        pool("void_eye_pool", "void_eye_stun", 45, forbid="TURN_MOD:5=0", cooldown=2),
        pool("dice_demon_pool", "dice_demon_fog", 36, cooldown=3),
        pool("dice_demon_pool", "dice_demon_control", 28, cooldown=2),
        pool("dice_demon_pool", "dice_demon_throw", 36),
    ]

    intents = [
        intent("goblin_snot", "BOARD", "鼻涕粘液", "生成鼻涕液", "生成一个鼻涕液地块", 2),
        intent("goblin_mace", "ATTACK", "狼牙棒", "攻击 9", "造成9点伤害", 1),
        intent("goblin_shield", "DEFEND", "盾牌", "护甲 6", "获得6点护甲", 1),
        intent("slime_attack", "ATTACK", "攻击", "攻击 10", "造成10点伤害", 1),
        intent("slime_defend", "DEFEND", "防御", "护甲 5", "获得5点护甲", 1),
        intent("slime_mud", "BOARD", "淤泥", "塞入淤泥", "塞入两张淤泥地块", 2),
        intent("skeleton_swing", "ATTACK", "骨头挥击", "攻击 11", "造成11点伤害", 1),
        intent("skeleton_spikes", "ATTACK", "骨刺", "攻击 3x3", "造成3x3点伤害", 2),
        intent("murloc_swamp_charge", "ATTACK", "沼泽突击", "攻击 1x6", "造成1x6点伤害", 2),
        intent("murloc_dive", "SPECIAL", "淤泥之下", "潜伏", "本回合免疫伤害，行动时力量+1", 2, immune=True, role="CONTROL"),
        intent("skeleton_archer_volley", "ATTACK", "齐射", "攻击 3x3", "造成3x3点伤害", 2),
        intent("skeleton_archer_empower", "SPECIAL", "强化", "力量 +1", "获得1点力量", 1, role="SCALING"),
        intent("frost_giant_ice_power", "SPECIAL", "冰霜之力", "骰子 1-3", "玩家3回合内骰子只会在1-3之间随机", 3),
        intent("frost_giant_ice_shield", "DEFEND", "冰盾", "护甲 9", "获得9点护甲", 1),
        intent("frost_giant_frost_strike", "ATTACK", "冰霜重击", "攻击 12", "造成12点伤害", 2),
        intent("guard_thrust", "ATTACK", "前刺", "攻击 11", "造成11点伤害", 1),
        intent("guard_formation", "DEFEND", "列阵", "护甲 8", "获得8点护盾", 1),
        intent("guard_bash", "ATTACK", "盾击", "攻击 7 护甲 6", "造成7点伤害并获得6点护盾", 2, role="MIXED"),
        intent("baby_dragon_flame", "ATTACK", "喷射火焰", "攻击 16", "造成16点伤害", 2),
        intent("baby_dragon_burn", "BOARD", "焚烧", "销毁地块", "随机销毁场上两个非诅咒地块", 3),
        intent("baby_dragon_tail", "ATTACK", "甩尾", "攻击 3x4", "造成3x4点伤害", 2),
        intent("skeleton_king_bone_spikes", "ATTACK", "骨刺", "骨头爆发", "造成(3+骨头数量)x3点伤害并消耗所有骨头", 3),
        intent("skeleton_king_bone_shield", "DEFEND", "骨盾", "护甲 8", "获得8点护盾", 1),
        intent("skeleton_king_slam", "ATTACK", "重击", "攻击 14", "造成14点伤害", 2),
        intent("werewolf_moon", "BOARD", "月圆之夜", "塞入月圆", "塞入月圆地块", 3),
        intent("werewolf_claw", "ATTACK", "爪击", "攻击 12", "造成12点伤害", 2),
        intent("werewolf_lick", "DEFEND", "舔舐伤口", "恢复 6", "恢复6点生命", 1),
        intent("werewolf_hysteria", "ATTACK", "歇斯底里", "攻击 24", "造成24点伤害", 3),
        intent("werewolf_moonbath", "DEFEND", "沐浴月光", "恢复 10", "恢复10点生命", 2),
        intent("headless_swing", "ATTACK", "挥击", "攻击 22", "造成22点伤害", 2),
        intent("headless_flame_combo", "ATTACK", "连击", "烈火追击", "每个烈火地块造成6点伤害", 2),
        intent("headless_wildfire", "BOARD", "野火", "生成烈火", "给玩家生成2张烈火地块", 2),
        intent("medusa_petrify", "SPECIAL", "石化凝视", "封锁骰子", "随机封锁玩家两个骰子", 3),
        intent("medusa_split_arrow", "BOARD", "分裂箭", "塞入分裂箭", "塞入三个分裂箭地块", 3),
        intent("medusa_snake_bite", "ATTACK", "蛇咬", "攻击 2x6", "造成2x6点伤害", 2),
        intent("void_eye_hot_ray", "ATTACK", "炙热射线", "攻击 2x5", "造成2x5点伤害", 2),
        intent("void_eye_void_gaze", "BOARD", "虚空凝视", "虚空化", "随机将非诅咒地块变为虚空地块，虚空大眼力量+1", 3),
        intent("void_eye_stun", "BOARD", "晕眩", "塞入晕眩", "塞入3张晕眩地块", 2),
        intent("dice_demon_fog", "SPECIAL", "嬉戏", "骰子迷雾", "玩家3回合无法得知骰子点数且不显示目标线", 3),
        intent("dice_demon_control", "SPECIAL", "掌控之力", "骰子固定1", "下一回合骰子只能骰出1点", 3),
        intent("dice_demon_throw", "ATTACK", "掷击", "攻击 2x6", "造成(2+力量)x6点伤害", 2),
    ]

    effects = [
        effect("goblin_snot", 1, "ADD_TILE", 1, "PLAYER_BOARD", snot_tile),
        effect("goblin_mace", 1, "DAMAGE", 9, "PLAYER"),
        effect("goblin_shield", 1, "BLOCK", 6, "SELF"),
        effect("slime_attack", 1, "DAMAGE", 10, "PLAYER"),
        effect("slime_defend", 1, "BLOCK", 5, "SELF"),
        effect("slime_mud", 1, "ADD_TILE", 2, "PLAYER_BOARD", mud_tile),
        effect("skeleton_swing", 1, "DAMAGE", 11, "PLAYER"),
        effect("skeleton_spikes", 1, "DAMAGE", 3, "PLAYER", hits=3),
        effect("murloc_swamp_charge", 1, "DAMAGE", 1, "PLAYER", hits=6),
        effect("murloc_dive", 1, "STRENGTH", 1, "SELF"),
        effect("skeleton_archer_volley", 1, "DAMAGE", 3, "PLAYER", hits=3),
        effect("skeleton_archer_empower", 1, "STRENGTH", 1, "SELF"),
        effect("frost_giant_ice_power", 1, "SET_DICE_RANGE", 3, "PLAYER_NEXT_TURNS", "1-3", duration=3),
        effect("frost_giant_ice_shield", 1, "BLOCK", 9, "SELF"),
        effect("frost_giant_frost_strike", 1, "DAMAGE", 12, "PLAYER"),
        effect("guard_thrust", 1, "DAMAGE", 11, "PLAYER"),
        effect("guard_formation", 1, "BLOCK", 8, "SELF"),
        effect("guard_bash", 1, "DAMAGE", 7, "PLAYER"),
        effect("guard_bash", 2, "BLOCK", 6, "SELF"),
        effect("baby_dragon_flame", 1, "DAMAGE", 16, "PLAYER"),
        effect("baby_dragon_burn", 1, "DESTROY_TILES", 2, "PLAYER_BOARD", "non_curse"),
        effect("baby_dragon_tail", 1, "DAMAGE", 3, "PLAYER", hits=4),
        effect("skeleton_king_bone_spikes", 1, "DAMAGE_BY_TILE_COUNT", 1, "PLAYER", bone_tile, base=3, hits=3),
        effect("skeleton_king_bone_spikes", 2, "DESTROY_TILE_ID", 0, "PLAYER_BOARD", bone_tile),
        effect("skeleton_king_bone_shield", 1, "BLOCK", 8, "SELF"),
        effect("skeleton_king_slam", 1, "DAMAGE", 14, "PLAYER"),
        effect("werewolf_moon", 1, "ADD_TILE", 1, "PLAYER_BOARD", moon_tile),
        effect("werewolf_claw", 1, "DAMAGE", 12, "PLAYER"),
        effect("werewolf_lick", 1, "HEAL", 6, "SELF"),
        effect("werewolf_hysteria", 1, "DAMAGE", 24, "PLAYER"),
        effect("werewolf_moonbath", 1, "HEAL", 10, "SELF"),
        effect("headless_swing", 1, "DAMAGE", 22, "PLAYER"),
        effect("headless_flame_combo", 1, "DAMAGE_PER_TILE", 6, "PLAYER", fire_tile),
        effect("headless_wildfire", 1, "ADD_TILE", 2, "PLAYER_BOARD", fire_tile),
        effect("medusa_petrify", 1, "LOCK_DICE", 2, "PLAYER_NEXT_TURN"),
        effect("medusa_split_arrow", 1, "ADD_TILE", 3, "PLAYER_BOARD", split_arrow_tile),
        effect("medusa_snake_bite", 1, "DAMAGE", 2, "PLAYER", hits=6),
        effect("void_eye_hot_ray", 1, "DAMAGE", 2, "PLAYER", hits=5),
        effect("void_eye_void_gaze", 1, "CORRUPT_TILE", 1, "RANDOM_PLAYER_TILE", void_tile),
        effect("void_eye_void_gaze", 2, "STRENGTH", 1, "SELF"),
        effect("void_eye_stun", 1, "ADD_TILE", 3, "PLAYER_BOARD", stun_tile),
        effect("dice_demon_fog", 1, "SET_DICE_FOG", 3, "PLAYER_NEXT_TURNS", duration=3),
        effect("dice_demon_control", 1, "SET_DICE_RANGE", 1, "PLAYER_NEXT_TURNS", "1-1", duration=1),
        effect("dice_demon_throw", 1, "DAMAGE", 2, "PLAYER", hits=6),
    ]

    config = {
        "battles": [
            {"battle": 1, "monster_id": "goblin", "rolls": 3},
            {"battle": 2, "monster_id": "slime", "rolls": 3},
            {"battle": 3, "monster_id": "skeleton_soldier", "rolls": 3},
            {"battle": 4, "monster_id": "murloc", "rolls": 3},
            {"battle": 5, "monster_id": "skeleton_archer", "rolls": 3},
            {"battle": 6, "monster_id": "frost_giant", "rolls": 3},
            {"battle": 7, "monster_id": "skeleton_pair", "rolls": 3},
            {"battle": 8, "monster_id": "guard_pair", "rolls": 3},
            {"battle": 9, "monster_id": "baby_dragon", "rolls": 3},
            {"battle": 10, "monster_id": "skeleton_king", "rolls": 3},
            {"battle": 11, "monster_id": "werewolf", "rolls": 3},
            {"battle": 12, "monster_id": "headless_knight", "rolls": 3},
            {"battle": 13, "monster_id": "medusa", "rolls": 3},
            {"battle": 14, "monster_id": "void_eye", "rolls": 3},
            {"battle": 15, "monster_id": "dice_demon", "rolls": 3},
        ],
        "monsters": monsters,
        "phases": phases,
        "intent_pools": pools,
        "intents": intents,
        "intent_effects": effects,
        "status_definitions": [],
    }
    write_json(MONSTER_CONFIG, config)


def sync_map() -> None:
    config = json.loads(MAP_CONFIG.read_text(encoding="utf-8"))
    config["monster_pool_entries"] = [
        {"pool_id": "act1_easy_normal", "monster_id": "goblin", "weight": 34},
        {"pool_id": "act1_easy_normal", "monster_id": "slime", "weight": 32},
        {"pool_id": "act1_easy_normal", "monster_id": "skeleton_soldier", "weight": 28},
        {"pool_id": "act1_mid_normal", "monster_id": "skeleton_archer", "weight": 26},
        {"pool_id": "act1_mid_normal", "monster_id": "murloc", "weight": 24},
        {"pool_id": "act1_mid_normal", "monster_id": "frost_giant", "weight": 18},
        {"pool_id": "act1_mid_normal", "monster_id": "goblin", "weight": 12},
        {"pool_id": "act1_late_normal", "monster_id": "skeleton_pair", "weight": 28},
        {"pool_id": "act1_late_normal", "monster_id": "guard_pair", "weight": 28},
        {"pool_id": "act1_late_normal", "monster_id": "baby_dragon", "weight": 24},
        {"pool_id": "act1_late_normal", "monster_id": "frost_giant", "weight": 16},
        {"pool_id": "act1_elite", "monster_id": "skeleton_king", "weight": 32},
        {"pool_id": "act1_elite", "monster_id": "werewolf", "weight": 30},
        {"pool_id": "act1_elite", "monster_id": "headless_knight", "weight": 30},
        {"pool_id": "act1_boss", "monster_id": "medusa", "weight": 34},
        {"pool_id": "act1_boss", "monster_id": "void_eye", "weight": 34},
        {"pool_id": "act1_boss", "monster_id": "dice_demon", "weight": 32},
    ]
    write_json(MAP_CONFIG, config)


def main() -> None:
    sync_tiles()
    sync_relics()
    sync_monsters()
    sync_map()


if __name__ == "__main__":
    main()
